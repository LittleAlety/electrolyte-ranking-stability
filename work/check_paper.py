# -*- coding: utf-8 -*-
"""核对结题论文是否满足 outputs/week35/paper_build_handoff.json 的 required_edits。

渲染约定：生成器把 _xxx 渲染为下标，因此正文里 N_resolved_both 会以 N+下标 出现，
词形检查同时接受下划线原形与去下划线形式。
"""
import io, re, sys
from docx import Document

DOCX = sys.argv[1]
SRC = sys.argv[2]
doc = Document(DOCX)
paras = [p.text for p in doc.paragraphs]
tbl = []
for t in doc.tables:
    for row in t.rows:
        tbl.append(" | ".join(c.text for c in row.cells))
text = "\n".join(paras) + "\n" + "\n".join(tbl)
flat = re.sub(r"[\s_]", "", text)
src = io.open(SRC, encoding="utf-8").read()

res = []


def chk(name, ok, detail):
    res.append((name, bool(ok), detail))


want = ["3.1  模型层级与外部参考边界", "3.2  电子结构与介质物理如何改变排序", "3.3  Li+ 配位条件态与还原态身份",
        "3.4  不确定度感知的材料筛选", "3.5  模型位移与配位修正的可预测性", "3.6  最小昂贵信息预算"]
pos = [text.find(w) for w in want]
chk("six_results_sections_in_fixed_order", all(p >= 0 for p in pos) and pos == sorted(pos), "positions=%s" % pos)

miss = [g for g in ["G%02d" % i for i in range(1, 14)] if g not in text]
chk("thirteen_boundary_statements_present", not miss, "missing=%s" % (miss or "none"))

miss = [l for l in ["L%02d" % i for i in range(1, 17)] if l not in text]
chk("sixteen_lineage_ids_present", not miss, "missing=%s" % (miss or "none"))

vals = ["0.4286", "15 / 6", "31", "0.6732", "0.7879", "0.7071068", "288", "296", "5920", "299", "918"]
bad = [v for v in vals if v not in text]
chk("lineage_values_present", not bad, "missing=%s" % (bad or "none"))

low = text.lower()
chk("designated_wording_used",
    "designated computational target" in low and "designated reference model" in low,
    "counts: target=%d model=%d" % (low.count("designated computational target"),
                                    low.count("designated reference model")))
chk("no_forbidden_wording",
    "validated target" not in low and "physically validated target" not in low,
    "forbidden occurrences: %d" % (low.count("validated target") + low.count("physically validated target")))

chk("T6a_T6b_definitional_and_sigma_convention",
    "T6a" in text and "T6b" in text and text.count("定义性") >= 3
    and "z \u2264 1/\u221a2" in text and "ddof" in text
    and "resolvedboth" in flat and "ndiscordantboth" in flat and "0.7071067811865475" in text,
    "T6a=%d T6b=%d 定义性=%d ddof=%d numerator/denominator stated=%s"
    % (text.count("T6a"), text.count("T6b"), text.count("定义性"), text.count("ddof"),
       ("resolvedboth" in flat)))

chk("F57_F58_placed_as_figures",
    "F57_metric_robustness.png" in src and "F58_layer_independence.png" in src
    and len(doc.inline_shapes) == 24 and len(doc.tables) == 19,
    "inline_shapes=%d (v6=22), tables=%d (v6=17)" % (len(doc.inline_shapes), len(doc.tables)))

chk("scope_limit_in_abstract_and_conclusion",
    "不是对候选分子绝对性能的排名" in text and "不是对候选分子真实电化学性能的排名" in text
    and low.count("designated computational target") >= 3,
    "scope clauses found")

for h in ["1  引言", "2  计算方法与决策量定义", "3  结果与讨论", "4  结语", "5  与预注册方案的对照",
          "6  结论边界、数字溯源与局限", "7  补充材料"]:
    chk("section_%s" % h.split()[0], h in text, h)

n_fail = sum(1 for _, ok, _ in res if not ok)
for name, ok, detail in res:
    print("%-46s %s  %s" % (name, "PASS" if ok else "FAIL", detail))
print("SUMMARY %d/%d passed" % (len(res) - n_fail, len(res)))
sys.exit(1 if n_fail else 0)
